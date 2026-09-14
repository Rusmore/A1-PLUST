"""Generate a PowerPoint overview of the Petty Cash Fund (PCF) Portal setup."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- Palette ---------------------------------------------------------------
INK      = RGBColor(0x1E, 0x29, 0x3B)   # deep navy text
BRAND    = RGBColor(0x0E, 0x7C, 0x66)   # teal accent
BRAND_DK = RGBColor(0x0A, 0x4A, 0x3E)
CREAM    = RGBColor(0xF6, 0xF4, 0xEF)   # app background
CARD     = RGBColor(0xFF, 0xFF, 0xFF)
MUT      = RGBColor(0x5B, 0x63, 0x72)
LINE     = RGBColor(0xDD, 0xD7, 0xCC)
AMBER    = RGBColor(0xB4, 0x7E, 0x0B)
RED      = RGBColor(0x8A, 0x10, 0x20)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide(bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    s.shapes._spTree.remove(r._element); s.shapes._spTree.insert(2, r._element)
    return s


def txt(slide, left, top, width, height, text, size=18, color=INK, bold=False,
        align=PP_ALIGN.LEFT, font="Segoe UI", anchor=MSO_ANCHOR.TOP, line_spacing=1.0):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        r = p.add_run(); r.text = ln
        f = r.font; f.size = Pt(size); f.bold = bold; f.name = font
        f.color.rgb = color
    return tb


def bar(slide, top, height, color=BRAND, left=0, width=SLIDE_W):
    b = slide.shapes.add_shape(1, left, top, width, height)
    b.fill.solid(); b.fill.fore_color.rgb = color; b.line.fill.background()
    b.shadow.inherit = False
    return b


def card(slide, left, top, width, height, fill=CARD, line=LINE):
    c = slide.shapes.add_shape(5, left, top, width, height)  # rounded rect
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.line.color.rgb = line; c.line.width = Pt(1)
    c.shadow.inherit = False
    return c


def bullets(slide, left, top, width, height, items, size=16, gap=6,
            color=INK, marker="•  ", bold_lead=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        lead, rest = (it.split("||", 1) + [""])[:2] if "||" in it else (None, it)
        if lead is not None:
            r0 = p.add_run(); r0.text = marker + lead + " "
            r0.font.size = Pt(size); r0.font.bold = True; r0.font.name = "Segoe UI"; r0.font.color.rgb = BRAND_DK
            r1 = p.add_run(); r1.text = rest
            r1.font.size = Pt(size); r1.font.name = "Segoe UI"; r1.font.color.rgb = color
        else:
            r = p.add_run(); r.text = marker + it
            r.font.size = Pt(size); r.font.bold = bold_lead; r.font.name = "Segoe UI"; r.font.color.rgb = color
    return tb


def header(slide, kicker, title):
    bar(slide, Inches(0), Inches(0), Inches(0.22), color=BRAND, width=SLIDE_W)
    txt(slide, Inches(0.6), Inches(0.42), Inches(12), Inches(0.4),
        kicker.upper(), size=13, color=BRAND, bold=True)
    txt(slide, Inches(0.6), Inches(0.75), Inches(12.1), Inches(0.9),
        title, size=30, color=INK, bold=True)
    bar(slide, Inches(0.62), Inches(1.62), Inches(0.03), color=LINE, width=Inches(11.9))


# ========================================================================
# 1. TITLE
# ========================================================================
s = add_slide(INK)
bar(s, Inches(0), Inches(0), Inches(0.28), color=BRAND, width=SLIDE_W)
bar(s, Inches(0), SLIDE_H - Inches(0.28), Inches(0.28), color=BRAND, width=SLIDE_W)
txt(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(0.5),
    "PETTY CASH FUND (PCF) PORTAL", size=16, color=BRAND, bold=True)
txt(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(1.5),
    "Petty Cash Management System", size=46, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
txt(s, Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.8),
    "System Overview, Architecture, Security & Hosting", size=20,
    color=RGBColor(0xC9, 0xD3, 0xD0))
txt(s, Inches(0.9), Inches(6.35), Inches(11.5), Inches(0.5),
    "Internal Finance Tool  ·  A1+ Group  ·  Prepared 2026", size=13,
    color=RGBColor(0x8A, 0x97, 0x9C))

# ========================================================================
# 2. WHAT IT IS
# ========================================================================
s = add_slide()
header(s, "Overview", "What the Portal Does")
txt(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.8),
    "A web-based system that replaces spreadsheet-based petty cash tracking with a single, "
    "auditable app across multiple companies, plants, and branches.", size=17, color=MUT)
lifecycle = ["Request", "Disburse", "Liquidate", "Replenish", "Monitor & Report"]
cw = Inches(2.28); gap = Inches(0.15); x = Inches(0.62); y = Inches(2.9)
for i, step in enumerate(lifecycle):
    card(s, x, y, cw, Inches(1.0))
    txt(s, x, y + Inches(0.30), cw, Inches(0.5), step, size=15, color=BRAND_DK,
        bold=True, align=PP_ALIGN.CENTER)
    x = Emu(x + cw + gap)
bullets(s, Inches(0.62), Inches(4.35), Inches(12), Inches(2.6), [
    "Full petty-cash lifecycle:||requests, disbursements, liquidation, replenishment, and aging.",
    "Role-based access:||each user sees only the plants/entities they are authorized for.",
    "Print-ready management reports||plus Excel/CSV export and an Acumatica Purchase Orders template.",
    "Complete audit trail||of every action for accountability.",
], size=16, gap=10)

# ========================================================================
# 3. KEY FEATURES
# ========================================================================
s = add_slide()
header(s, "Capabilities", "Key Features")
feats = [
    ("Dashboard", "KPI cards and real-time fund monitoring per plant."),
    ("Requests & Disbursements", "Track approvals and maintain a release ledger."),
    ("Liquidation", "Receipt-level amounts, reconciliation, cash settlement."),
    ("Replenishment", "Restore funds back to their float amount."),
    ("Liquidation Aging", "Monitor unliquidated balances over time."),
    ("Report Center", "Print/PDF output + Excel/CSV builders."),
    ("Master Data", "Branches, companies, plants, tax & expense categories."),
    ("Audit Trail", "Immutable-style log of user actions."),
]
cw = Inches(5.95); ch = Inches(1.15); gx = Inches(0.3); gy = Inches(0.25)
x0 = Inches(0.62); y0 = Inches(1.85)
for i, (t, d) in enumerate(feats):
    col = i % 2; row = i // 2
    x = Emu(x0 + col * (cw + gx)); y = Emu(y0 + row * (ch + gy))
    card(s, x, y, cw, ch)
    bar(s, x, y, ch, color=BRAND, width=Inches(0.08))
    txt(s, Emu(x + Inches(0.28)), Emu(y + Inches(0.14)), Emu(cw - Inches(0.4)), Inches(0.4),
        t, size=16, color=BRAND_DK, bold=True)
    txt(s, Emu(x + Inches(0.28)), Emu(y + Inches(0.55)), Emu(cw - Inches(0.4)), Inches(0.5),
        d, size=13, color=MUT)

# ========================================================================
# 4. ARCHITECTURE
# ========================================================================
s = add_slide()
header(s, "How It Works", "Architecture — No Build Step")
txt(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(0.7),
    "A single-page React app that runs directly in the browser. No bundler, no npm install.",
    size=17, color=MUT)
steps = [
    ("index.html", "Loads Babel + an importmap (react, recharts, xlsx, lucide-react)."),
    ("Loader", "Fetches the ordered src/*.jsx fragments and joins them in order."),
    ("Babel", "Transpiles the combined code once, in the browser."),
    ("Run", "Executes as one ES module → the app mounts."),
]
x = Inches(0.62); y = Inches(2.75); cw = Inches(2.9); ch = Inches(1.6)
for i, (t, d) in enumerate(steps):
    card(s, x, y, cw, ch)
    txt(s, x, Emu(y + Inches(0.18)), cw, Inches(0.4), f"{i+1}", size=22, color=BRAND, bold=True, align=PP_ALIGN.CENTER)
    txt(s, x, Emu(y + Inches(0.6)), cw, Inches(0.4), t, size=15, color=INK, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Emu(x + Inches(0.15)), Emu(y + Inches(0.95)), Emu(cw - Inches(0.3)), Inches(0.6),
        d, size=11.5, color=MUT, align=PP_ALIGN.CENTER)
    if i < 3:
        txt(s, Emu(x + cw - Inches(0.05)), Emu(y + Inches(0.55)), Inches(0.4), Inches(0.5),
            "›", size=30, color=BRAND, bold=True, align=PP_ALIGN.CENTER)
    x = Emu(x + cw + Inches(0.18))
card(s, Inches(0.62), Inches(4.7), Inches(12.1), Inches(1.9), fill=RGBColor(0xEF, 0xF5, 0xF3))
txt(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(0.4), "The source is split into ~22 ordered fragments (src/)", size=15, color=BRAND_DK, bold=True)
bullets(s, Inches(0.85), Inches(5.35), Inches(11.6), Inches(1.2), [
    "Files share one scope; order matters (later files depend on earlier ones).",
    "Bump PCP_SRC_VERSION on each deploy to bust the browser cache.",
    "Tech stack: React 18, Recharts, SheetJS (xlsx), lucide-react, optional Supabase.",
], size=14, gap=5)

# ========================================================================
# 5. DATA STORAGE
# ========================================================================
s = add_slide()
header(s, "Where Data Lives", "Data Storage")
cols = [
    ("Supabase", "Primary / shared", [
        "One JSON blob in the pcp_state table",
        "Source of truth across all devices",
        "Debounced saves (~800ms)",
    ], BRAND),
    ("Browser (localStorage)", "Local cache", [
        "Instant response + offline copy",
        "Per-browser mirror of the cloud",
    ], MUT),
    ("Backups", "Snapshots", [
        "Periodic in-app state snapshots",
        "Recovery path if a write goes wrong",
    ], BRAND_DK),
]
x = Inches(0.62); cw = Inches(3.95); ch = Inches(2.9); y = Inches(1.95)
for t, sub, items, clr in cols:
    card(s, x, y, cw, ch)
    bar(s, x, y, Inches(0.5), color=clr, width=cw)
    txt(s, Emu(x + Inches(0.25)), Emu(y + Inches(0.62)), Emu(cw - Inches(0.4)), Inches(0.4), t, size=17, color=INK, bold=True)
    txt(s, Emu(x + Inches(0.25)), Emu(y + Inches(1.0)), Emu(cw - Inches(0.4)), Inches(0.35), sub, size=12.5, color=clr, bold=True)
    bullets(s, Emu(x + Inches(0.25)), Emu(y + Inches(1.45)), Emu(cw - Inches(0.45)), Inches(1.3), items, size=12.5, gap=6)
    x = Emu(x + cw + Inches(0.2))
card(s, Inches(0.62), Inches(5.15), Inches(12.1), Inches(1.5), fill=RGBColor(0xFD, 0xF6, 0xE9), line=RGBColor(0xE7, 0xD5, 0xA8))
txt(s, Inches(0.85), Inches(5.32), Inches(11.6), Inches(0.4), "⚠  Uploaded files are embedded as base64 inside the same JSON blob", size=15, color=AMBER, bold=True)
txt(s, Inches(0.85), Inches(5.75), Inches(11.6), Inches(0.8),
    "Receipts, PCF documents and reimbursement attachments are stored as text in the state — not as real files. "
    "This inflates size ~33% and counts against the 500 MB free tier.", size=13, color=MUT)

# ========================================================================
# 6. SECURITY & ACCESS
# ========================================================================
s = add_slide()
header(s, "Security", "Access Control & Security")
card(s, Inches(0.62), Inches(1.85), Inches(5.95), Inches(4.6), fill=RGBColor(0xEF, 0xF5, 0xF3))
txt(s, Inches(0.85), Inches(2.0), Inches(5.5), Inches(0.4), "In place ✓", size=17, color=BRAND_DK, bold=True)
bullets(s, Inches(0.85), Inches(2.5), Inches(5.5), Inches(3.9), [
    "Row Level Security enabled on pcp_state",
    "Authenticated-only read / insert / update policies",
    "Public sign-ups disabled — admin adds users",
    "Anon key safe to expose (protected by RLS)",
    "Role + plant-based access per user",
    "Delete reserved for SuperAdmin",
    "Single designated Liquidation Approver",
    "HTTPS enforced by the static host",
], size=14.5, gap=9)
card(s, Inches(6.78), Inches(1.85), Inches(5.95), Inches(4.6), fill=RGBColor(0xFC, 0xF1, 0xF2))
txt(s, Inches(7.0), Inches(2.0), Inches(5.5), Inches(0.4), "Watch / To harden", size=17, color=RED, bold=True)
bullets(s, Inches(7.0), Inches(2.5), Inches(5.5), Inches(3.9), [
    "RLS checks 'authenticated', not role/entity",
    "Single-blob model → any signed-in user can write all data via the API",
    "Role limits are enforced in the app (client) layer only",
    "Make the GitHub repo private (emails/source are exposed if public)",
    "True per-role/per-entity isolation needs the DB refactor (tech-debt #2)",
], size=14.5, gap=10, color=INK)

# ========================================================================
# 7. ROLES  (incl. new PCF Requestor)
# ========================================================================
s = add_slide()
header(s, "Roles", "User Roles & the New PCF Requestor")
rows = [
    ("Role", "Access", True),
    ("SuperAdmin / Accounting", "Full system access, all plants, user & settings admin", False),
    ("Finance", "All modules across plants (no user admin)", False),
    ("Custodian", "Own plant(s): requests, liquidation, replenishment, reports", False),
    ("PCF Requestor  (NEW)", "Prepares transactions for ONE entity only", False),
]
x = Inches(0.62); y = Inches(1.85); w1 = Inches(4.3); w2 = Inches(7.8); rh = Inches(0.62)
for i, (a, b, hd) in enumerate(rows):
    fill = BRAND if hd else (RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF0, 0xEC, 0xE3))
    c1 = s.shapes.add_shape(1, x, y, w1, rh); c1.fill.solid(); c1.fill.fore_color.rgb = fill; c1.line.color.rgb = LINE; c1.line.width = Pt(0.5); c1.shadow.inherit = False
    c2 = s.shapes.add_shape(1, Emu(x + w1), y, w2, rh); c2.fill.solid(); c2.fill.fore_color.rgb = fill; c2.line.color.rgb = LINE; c2.line.width = Pt(0.5); c2.shadow.inherit = False
    tc = RGBColor(0xFF, 0xFF, 0xFF) if hd else INK
    txt(s, Emu(x + Inches(0.2)), y, Emu(w1 - Inches(0.3)), rh, a, size=14, color=(tc if not hd else tc), bold=(hd or "NEW" in a), anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Emu(x + w1 + Inches(0.2)), y, Emu(w2 - Inches(0.4)), rh, b, size=13.5, color=tc, bold=hd, anchor=MSO_ANCHOR.MIDDLE)
    y = Emu(y + rh)
card(s, Inches(0.62), Inches(5.05), Inches(12.1), Inches(1.65), fill=RGBColor(0xEF, 0xF5, 0xF3))
txt(s, Inches(0.85), Inches(5.2), Inches(11.6), Inches(0.4), "PCF Requestor permissions", size=15, color=BRAND_DK, bold=True)
bullets(s, Inches(0.85), Inches(5.62), Inches(11.6), Inches(1.0), [
    "Petty Cash Requests — full (create/edit/submit/upload); no approve/reject/release.",
    "Release Ledger — view only.   Liquidation — full except approval.   Data scoped to one entity.",
], size=13.5, gap=5)

# ========================================================================
# 8. HOSTING
# ========================================================================
s = add_slide()
header(s, "Deployment", "Hosting Setup")
txt(s, Inches(0.6), Inches(1.75), Inches(12.1), Inches(0.6),
    "100% static files — any static host works. GitHub is the repository; a static host serves the site.",
    size=16, color=MUT)
card(s, Inches(0.62), Inches(2.5), Inches(5.95), Inches(2.0))
txt(s, Inches(0.85), Inches(2.65), Inches(5.5), Inches(0.4), "GitHub — repository", size=16, color=BRAND_DK, bold=True)
bullets(s, Inches(0.85), Inches(3.1), Inches(5.5), Inches(1.3), [
    "Source of truth & version history",
    "Recommend: make the repo PRIVATE",
    "Auto-deploys on push",
], size=13.5, gap=7)
card(s, Inches(6.78), Inches(2.5), Inches(5.95), Inches(2.0))
txt(s, Inches(7.0), Inches(2.65), Inches(5.5), Inches(0.4), "Host — serves the site", size=16, color=BRAND_DK, bold=True)
bullets(s, Inches(7.0), Inches(3.1), Inches(5.5), Inches(1.3), [
    "Cloudflare Pages (recommended, free, generous)",
    "or GitHub Pages / Netlify / Vercel",
    "Free HTTPS + custom domain",
], size=13.5, gap=7)
card(s, Inches(0.62), Inches(4.75), Inches(12.1), Inches(1.9), fill=RGBColor(0xFD, 0xF6, 0xE9), line=RGBColor(0xE7, 0xD5, 0xA8))
txt(s, Inches(0.85), Inches(4.9), Inches(11.6), Inches(0.4), "When switching hosts / domains", size=15, color=AMBER, bold=True)
bullets(s, Inches(0.85), Inches(5.35), Inches(11.6), Inches(1.2), [
    "Update Supabase → Authentication → URL Configuration (Site URL + Redirect URLs), or password resets break.",
    "Build command: none.   Output directory: root ( / ).   No code changes needed.",
    "Cost: $0 at your scale — the only possible future cost is Supabase Pro (~$25/mo) for backups/space.",
], size=13.5, gap=6)

# ========================================================================
# 9. CAPACITY
# ========================================================================
s = add_slide()
header(s, "Scale", "Capacity — Can It Handle Daily Users?")
card(s, Inches(0.62), Inches(1.85), Inches(5.95), Inches(2.3), fill=RGBColor(0xEF, 0xF5, 0xF3))
txt(s, Inches(0.85), Inches(2.0), Inches(5.5), Inches(0.4), "Not the bottleneck ✓", size=16, color=BRAND_DK, bold=True)
bullets(s, Inches(0.85), Inches(2.5), Inches(5.5), Inches(1.6), [
    "Static host serves unlimited readers",
    "Supabase free tier >> 10 daily users",
    "Transpile cost is on the client, not server",
], size=14, gap=9)
card(s, Inches(6.78), Inches(1.85), Inches(5.95), Inches(2.3), fill=RGBColor(0xFC, 0xF1, 0xF2))
txt(s, Inches(7.0), Inches(2.0), Inches(5.5), Inches(0.4), "The real limit", size=16, color=RED, bold=True)
bullets(s, Inches(7.0), Inches(2.5), Inches(5.5), Inches(1.6), [
    "Single shared JSON blob, last-write-wins",
    "Simultaneous edits can overwrite silently",
    "Embedded files grow every save",
], size=14, gap=9)
card(s, Inches(0.62), Inches(4.35), Inches(12.1), Inches(2.2), fill=CARD)
txt(s, Inches(0.85), Inches(4.5), Inches(11.6), Inches(0.4), "Verdict", size=16, color=INK, bold=True)
bullets(s, Inches(0.85), Inches(4.95), Inches(11.6), Inches(1.5), [
    "Good for ~10 users doing mostly non-overlapping work||— which the per-entity accounts encourage.",
    "Risk appears with concurrent editing of the same data,||not with load volume.",
    "Scales in load; the fix for concurrency is the per-record table refactor.",
], size=14, gap=8)

# ========================================================================
# 10. TECH DEBT / RECOMMENDATIONS
# ========================================================================
s = add_slide()
header(s, "Roadmap", "Recommendations & Tech Debt")
items = [
    ("1", "Move uploaded files to Supabase Storage",
     "Stop embedding base64 in the state blob — protects the 500 MB free tier and keeps saves fast.", BRAND),
    ("2", "Split the state blob into per-record tables",
     "Ends last-write-wins data loss AND enables true DB/API-level role & entity isolation.", AMBER),
    ("3", "Automated Supabase backups + audit-log integrity",
     "Scheduled exports; move the audit trail to an append-only table.", BRAND_DK),
]
y = Inches(1.95)
for num, t, d, clr in items:
    card(s, Inches(0.62), y, Inches(12.1), Inches(1.35))
    bar(s, Inches(0.62), y, Inches(1.35), color=clr, width=Inches(0.12))
    txt(s, Inches(0.85), Emu(y + Inches(0.28)), Inches(0.9), Inches(0.8), num, size=34, color=clr, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1.85), Emu(y + Inches(0.2)), Inches(10.5), Inches(0.5), t, size=17, color=INK, bold=True)
    txt(s, Inches(1.85), Emu(y + Inches(0.68)), Inches(10.6), Inches(0.6), d, size=13.5, color=MUT)
    y = Emu(y + Inches(1.5))
txt(s, Inches(0.62), Inches(6.75), Inches(12), Inches(0.4),
    "Verdict: solid, production-ready MVP for a small team. Address #1 first as file volume grows.",
    size=13, color=BRAND_DK, bold=True)

# ========================================================================
# 11. NEXT STEPS
# ========================================================================
s = add_slide(INK)
bar(s, Inches(0), Inches(0), Inches(0.22), color=BRAND, width=SLIDE_W)
txt(s, Inches(0.8), Inches(0.7), Inches(11), Inches(0.8), "Immediate Next Steps", size=32,
    color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
steps = [
    "Confirm Supabase public sign-ups are OFF and the Users list is clean.",
    "Create the three PCF Requestor logins in Supabase (matching emails).",
    "Make the GitHub repository private.",
    "Commit & push pending changes (cleanup, Requestor role, docs).",
    "Verify the deploy in the Actions tab; hard-refresh and test a Requestor login.",
    "Plan tech-debt #1 (file storage) before receipt volume grows.",
]
y = Inches(1.9)
for i, st in enumerate(steps):
    n = s.shapes.add_shape(9, Inches(0.85), y, Inches(0.5), Inches(0.5))  # oval
    n.fill.solid(); n.fill.fore_color.rgb = BRAND; n.line.fill.background(); n.shadow.inherit = False
    txt(s, Inches(0.85), y, Inches(0.5), Inches(0.5), str(i + 1), size=18, color=RGBColor(0xFF,0xFF,0xFF), bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Inches(1.6), y, Inches(11), Inches(0.5), st, size=16, color=RGBColor(0xE7, 0xEC, 0xEA), anchor=MSO_ANCHOR.MIDDLE)
    y = Emu(y + Inches(0.78))

prs.save("PettyCashPortal_Overview.pptx")
print("Saved PettyCashPortal_Overview.pptx with", len(prs.slides._sldIdLst), "slides")
